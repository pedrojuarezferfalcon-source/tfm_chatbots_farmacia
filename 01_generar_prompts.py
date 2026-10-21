"""
01_generar_prompts.py
======================
TFM - Comparación de Chatbots clínicos en la práctica farmacéutica

Lee la plantilla de validación (hojas "Contraindicaciones" e "Interacciones")
y genera, para cada caso CONFIRMADO, el prompt estandarizado descrito en el
apartado 5.3.2 de la memoria:

  "Paciente de [edad] años, [sexo], diagnosticado de [condición].
   Medicación habitual: [fármaco A, dosis]. Se le prescribe [fármaco B, dosis].
   ¿Existe algún riesgo farmacológico relevante?
   ¿Qué recomendación haría el farmacéutico?"

LIMITACIÓN DOCUMENTADA (piloto intermedio):
La plantilla no incluye edad/sexo/condición reales por caso (ese dato vendría
de los metadatos demográficos de FAERS, no capturados en esta plantilla).
Para este piloto se genera una edad y un sexo PSEUDOALEATORIOS pero
REPRODUCIBLES por caso (semilla fija = id_caso, así que si vuelves a
ejecutar el script obtienes siempre los mismos prompts). No pretenden ser
demografía clínicamente relevante -edad/sexo no condicionan, en general, si
existe o no la contraindicación/interacción de cada caso- sino simplemente
dar variedad realista a los escenarios en vez de repetir el mismo patrón
55 años/sexo no especificado en los 35 prompts. La única excepción es que
el sexo se fija a "mujer" cuando la condición contraindicada es embarazo o
lactancia, para que el escenario sea clínicamente coherente. Esto debe
declararse como limitación metodológica en la entrega intermedia y
sustituirse por demografía representativa derivada de los metadatos FAERS
reales antes de la entrega final con los 100 casos.

La dosis de cada fármaco se deja como "dosis estándar" por el mismo motivo
(la plantilla no registra posología concreta).

Salidas:
  - prompts.json   -> lista de casos con su prompt, para el pipeline de llamadas
  - prompts_preview.xlsx -> para revisión humana rápida antes de gastar llamadas de API
"""

import json
import random
from pathlib import Path

import openpyxl

PLANTILLA = Path(__file__).parent / "plantilla_validacion_casos.xlsx"
OUT_JSON = Path(__file__).parent / "prompts.json"
OUT_PREVIEW = Path(__file__).parent / "prompts_preview.xlsx"

DOSIS_GENERICA = "dosis estándar"

# Rango de edad adulta plausible para el patrón demográfico variado (ver
# limitación arriba). No tiene relevancia clínica en sí, solo da variedad.
EDAD_MIN, EDAD_MAX = 22, 85
SEXOS = ["hombre", "mujer"]

# Condiciones en las que, por coherencia clínica mínima, el sexo se fuerza
# a "mujer" (no tiene sentido un escenario de embarazo en un hombre).
CONDICIONES_SOLO_MUJER = ("embarazo", "lactancia", "gestacion", "gestación")


def humanize(cond: str) -> str:
    """convierte snake_case tipo 'insuficiencia_renal' en texto legible"""
    return cond.replace("_", " ").strip()


def generar_demografia(id_caso: str, condicion: str | None = None) -> tuple[int, str]:
    """Genera (edad, sexo) pseudoaleatorios pero reproducibles para un caso.

    Semilla = id_caso, así que re-ejecutar el script da siempre los mismos
    prompts (importante para la reproducibilidad exigida en 5.5.1 de la
    memoria). No representa demografía real de FAERS: es solo variedad.
    """
    rng = random.Random(f"demografia-{id_caso}")
    edad = rng.randint(EDAD_MIN, EDAD_MAX)
    condicion_lower = (condicion or "").lower()
    if any(c in condicion_lower for c in CONDICIONES_SOLO_MUJER):
        sexo = "mujer"
    else:
        sexo = rng.choice(SEXOS)
    return edad, sexo


def construir_prompt_interaccion(id_caso, farmaco_a, farmaco_b) -> str:
    edad, sexo = generar_demografia(id_caso)
    return (
        f"Paciente de {edad} años, {sexo}, sin otras comorbilidades "
        f"relevantes conocidas. Medicación habitual: "
        f"{farmaco_a.title()} ({DOSIS_GENERICA}). Se le prescribe "
        f"{farmaco_b.title()} ({DOSIS_GENERICA}). ¿Existe algún riesgo "
        f"farmacológico relevante? ¿Qué recomendación haría el farmacéutico?"
    )


def construir_prompt_contraindicacion(id_caso, farmaco, condicion) -> str:
    edad, sexo = generar_demografia(id_caso, condicion)
    return (
        f"Paciente de {edad} años, {sexo}, diagnosticado de "
        f"{humanize(condicion)}. Se le prescribe {farmaco.title()} "
        f"({DOSIS_GENERICA}). ¿Existe algún riesgo farmacológico relevante? "
        f"¿Qué recomendación haría el farmacéutico?"
    )


def leer_contraindicaciones(wb):
    ws = wb["Contraindicaciones"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    casos = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx["nº_caso"]] is None:
            continue
        if (row[idx["estado"]] or "").strip().lower() != "confirmado":
            continue
        farmaco = row[idx["farmaco"]]
        condicion = row[idx["condicion_o_situacion_contraindicada"]]
        casos.append(
            {
                "id_caso": row[idx["nº_caso"]],
                "categoria": "Contraindicacion",
                "categoria_contraindicacion": row[idx["categoria_contraindicacion"]],
                "farmaco_a": farmaco,
                "farmaco_b": None,
                "condicion": condicion,
                "fuente_consultada": row[idx["fuente_consultada"]],
                "seccion_citada": str(row[idx["seccion_citada"]]),
                "url_o_referencia": row[idx["url_o_referencia"]] if "url_o_referencia" in idx else None,
                "ground_truth": row[idx["hallazgo_resumen"]]
                or f"{farmaco.title()} está contraindicado en {humanize(condicion)} "
                f"según ficha técnica ({row[idx['fuente_consultada']]}, "
                f"sección {row[idx['seccion_citada']]}).",
                "prompt": construir_prompt_contraindicacion(
                    row[idx["nº_caso"]], farmaco, condicion
                ),
            }
        )
    return casos


def leer_interacciones(wb):
    ws = wb["Interacciones"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    casos = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx["nº_caso"]] is None:
            continue
        if (row[idx["estado"]] or "").strip().lower() != "confirmado":
            continue
        farmaco_a = row[idx["farmaco_a"]]
        farmaco_b = row[idx["farmaco_b"]]
        casos.append(
            {
                "id_caso": row[idx["nº_caso"]],
                "categoria": "Interaccion",
                "severidad_mecanismo": row[idx["severidad_mecanismo"]],
                "farmaco_a": farmaco_a,
                "farmaco_b": farmaco_b,
                "condicion": None,
                "fuente_consultada": row[idx["fuente_consultada"]],
                "seccion_citada": row[idx["seccion_citada"]],
                "url_o_referencia": row[idx["url_o_referencia"]] if "url_o_referencia" in idx else None,
                "ground_truth": row[idx["hallazgo_resumen"]]
                or f"Interacción {row[idx['severidad_mecanismo']]} descrita entre "
                f"{farmaco_a.title()} y {farmaco_b.title()}.",
                "prompt": construir_prompt_interaccion(
                    row[idx["nº_caso"]], farmaco_a, farmaco_b
                ),
            }
        )
    return casos


def main():
    wb = openpyxl.load_workbook(PLANTILLA, data_only=True)
    casos = leer_contraindicaciones(wb) + leer_interacciones(wb)

    print(f"Casos confirmados encontrados: {len(casos)}")
    print(f"  - Contraindicaciones: {sum(1 for c in casos if c['categoria']=='Contraindicacion')}")
    print(f"  - Interacciones:      {sum(1 for c in casos if c['categoria']=='Interaccion')}")

    OUT_JSON.write_text(json.dumps(casos, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\nEscrito: {OUT_JSON}")

    # Vista previa en Excel para revisión humana rápida antes de llamar a las APIs
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "prompts_preview"
    cols = ["id_caso", "categoria", "prompt", "ground_truth", "fuente_consultada", "url_o_referencia"]
    ws_out.append(cols)
    for c in casos:
        ws_out.append([c.get(k, "") for k in cols])
    for col_cells in ws_out.columns:
        ws_out.column_dimensions[col_cells[0].column_letter].width = 45
    wb_out.save(OUT_PREVIEW)
    print(f"Escrito: {OUT_PREVIEW}  (revísalo antes de gastar llamadas de API)")


if __name__ == "__main__":
    main()
